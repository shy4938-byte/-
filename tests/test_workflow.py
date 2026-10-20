import unittest
import tempfile
from pathlib import Path
import fitz
from openpyxl import load_workbook
from speccompare.parser import analyze, AnalysisError, normalize
from speccompare.excel import export, inspect_template
from speccompare.schema import FIELDS, Product, Value
ROOT=Path(__file__).resolve().parents[1]
def pdf(path,lines):
    doc=fitz.open(); page=doc.new_page(); page.insert_text((40,50),'\n'.join(lines)); doc.save(path); doc.close()
class Workflow(unittest.TestCase):
    def setUp(self): self.tmp=tempfile.TemporaryDirectory(); self.root=Path(self.tmp.name)
    def tearDown(self): self.tmp.cleanup()
    def test_all_categories(self):
        cases={'MEMORY':['Manufacturer: Samsung','Model: SAMPLE-DDR5','Capacity: 16 GB','Memory Type: DDR5 UDIMM','Data Rate: 5600 MT/s','CAS Latency: CL46'], 'SSD':['Manufacturer: ADATA','Model: SAMPLE-SSD','Capacity: 1 TB','Sequential Read: 7.4 GB/s','Random Read: 1000000 IOPS','DRAM: Yes','TBW: 600 TBW'], 'VGA':['Manufacturer: NVIDIA','Model: SAMPLE-GPU','Memory Size: 16 GB','Memory Type: GDDR6','Memory Bandwidth: 512 GB/s','TGP: 200 W']}
        for category,lines in cases.items():
            with self.subTest(category=category):
                path=self.root/(category+'.pdf'); pdf(path,lines); p=analyze(path,category)[0]
                self.assertNotEqual(p.name,'확인 필요'); self.assertEqual(p.values['가격'].status,'미확인')
                self.assertTrue(p.values['제조사'].evidence); self.assertEqual(p.values['제조사'].page,1)
                if category=='SSD': self.assertEqual(p.values['순차읽기 속도'].text,'7400 MB/s')
                if category=='VGA': self.assertEqual(p.values['소비전력'].text,'200 W (TGP)')
    def test_no_model_mixing(self):
        path=self.root/'multiple.pdf'; pdf(path,['Model: A','Capacity: 1 TB','Model: B','Capacity: 2 TB'])
        p=analyze(path,'SSD')[0]; self.assertEqual(p.values['용량'].text,'확인 필요'); self.assertTrue(p.warnings)
    def test_model_table(self):
        path=self.root/'table.pdf'; doc=fitz.open(); page=doc.new_page()
        rows=[['Model','A','B'],['Capacity','1 TB','2 TB'],['Sequential Read','5000 MB/s','7000 MB/s']]
        xs=[40,210,380,550]; ys=[40,80,120,160]
        for x in xs: page.draw_line((x,40),(x,160))
        for y in ys: page.draw_line((40,y),(550,y))
        for r,row in enumerate(rows):
            for c,text in enumerate(row): page.insert_text((xs[c]+5,ys[r]+23),text)
        doc.save(path); doc.close()
        products=analyze(path,'SSD'); self.assertEqual(len(products),2)
        self.assertEqual(products[0].values['용량'].text,'1 TB'); self.assertEqual(products[1].values['순차읽기 속도'].text,'7000 MB/s')
    def test_invalid_scan_encryption(self):
        path=self.root/'bad.pdf'; path.write_text('bad')
        with self.assertRaises(AnalysisError): analyze(path,'MEMORY')
        doc=fitz.open(); doc.new_page(); doc.save(path); doc.close()
        with self.assertRaisesRegex(AnalysisError,'스캔'): analyze(path,'MEMORY')
        doc=fitz.open(); doc.new_page(); encrypted=self.root/'locked.pdf'; doc.save(encrypted,encryption=fitz.PDF_ENCRYPT_AES_256,owner_pw='owner',user_pw='secret'); doc.close()
        with self.assertRaisesRegex(AnalysisError,'암호화'): analyze(encrypted,'SSD')
    def test_template_preservation(self):
        template=ROOT/'assets/template.xlsx'; original=template.read_bytes()
        for category in FIELDS:
            with self.subTest(category=category):
                name,mapping=inspect_template(template,category)
                self.assertEqual(list(mapping.values()),list(range(5,5+len(FIELDS[category]))))
                p=Product('sample',category,{k:Value('=1+1') for k in FIELDS[category]})
                out=self.root/(category+'한글.xlsx'); export(template,out,category,[p])
                before,after=load_workbook(template),load_workbook(out)
                self.assertEqual(before.sheetnames,after.sheetnames)
                for s in before:
                    t=after[s.title]
                    self.assertEqual(str(s.merged_cells),str(t.merged_cells)); self.assertEqual(s.print_area,t.print_area); self.assertEqual(s.page_setup,t.page_setup)
                    self.assertEqual({k:(v.width,v.hidden,tuple(v._style or ()),v.min,v.max) for k,v in s.column_dimensions.items()},{k:(v.width,v.hidden,tuple(v._style or ()),v.min,v.max) for k,v in t.column_dimensions.items()}); self.assertEqual({k:(v.height,v.hidden,tuple(v._style or ())) for k,v in s.row_dimensions.items()},{k:(v.height,v.hidden,tuple(v._style or ())) for k,v in t.row_dimensions.items()})
                    for row in s:
                        for cell in row:
                            other=t[cell.coordinate]
                            from copy import copy
                            if s.title==name and cell.column in range(3,7) and cell.row in mapping.values():
                                self.assertEqual(other.alignment.horizontal,'center');self.assertEqual(other.alignment.vertical,'center')
                                self.assertEqual(copy(cell.font),copy(other.font));self.assertEqual(copy(cell.border),copy(other.border));self.assertEqual(copy(cell.fill),copy(other.fill))
                            else:self.assertEqual(cell._style,other._style)
                            if s.title!=name or cell.column not in range(3,7) or cell.row not in mapping.values(): self.assertEqual(cell.value,other.value)
                import zipfile
                with zipfile.ZipFile(template) as a, zipfile.ZipFile(out) as b:
                    changed=[n for n in a.namelist() if a.read(n)!=b.read(n)]
                    self.assertTrue(1<=len(changed)<=2); self.assertEqual(len([n for n in changed if n.startswith('xl/worksheets/')]),1); self.assertTrue(all(n.startswith('xl/worksheets/') or n=='xl/styles.xml' for n in changed))
                self.assertEqual(after[name]['C5'].data_type,'s'); before.close(); after.close()
        self.assertEqual(template.read_bytes(),original)
        with self.assertRaises(ValueError): export(template,template,'MEMORY',[])
    def test_units(self):
        self.assertEqual(normalize('7.4 GB/s'),'7400 MB/s'); self.assertEqual(normalize('3200 MHz'),'3200 MHz')
if __name__=='__main__': unittest.main()
