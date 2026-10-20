"""Integration test: set SPECCOMPARE_PM9E1_PDF to the user-supplied Rev 1.3 PDF."""
import os
import tempfile
import unittest
from pathlib import Path
from openpyxl import load_workbook
from speccompare.parser import analyze
from speccompare.excel import export
PDF=os.environ.get('SPECCOMPARE_PM9E1_PDF')
@unittest.skipUnless(PDF,'PM9E1 integration PDF not supplied')
class PM9E1(unittest.TestCase):
    def test_real_datasheet(self):
        products=analyze(PDF,'SSD')
        self.assertEqual([p.values['부품번호'].text for p in products],['MZVLC512HFJD-00B07','MZVLC1T0HFLU-00B07','MZVLC2T0HBLD-00B07','MZVLC4T0HBL1-00B07'])
        for i,p in enumerate(products):
            expected={'용량':['512GB','1024GB','2048GB','4096GB'][i], '순차읽기 속도':['12,200 MB/s','14,000 MB/s','14,000 MB/s','14,500 MB/s'][i], '순차쓰기 속도':['7,000 MB/s','12,500 MB/s','12,500 MB/s','12,600 MB/s'][i], '랜덤읽기 속도':['800K IOPS','1,600K IOPS','1,600K IOPS','2,000K IOPS'][i], '랜덤쓰기 속도':['1,360K IOPS','2,600K IOPS','2,600K IOPS','2,640K IOPS'][i], '내구성':['300 TB','600 TB','1,200 TB','2,400 TB'][i]}
            for key,value in expected.items():
                self.assertEqual(p.values[key].text,value); self.assertTrue(p.values[key].evidence)
                self.assertEqual(p.values[key].status,'검토 필요')
            self.assertEqual(p.values['가격'].text,'확인 필요')
        with tempfile.TemporaryDirectory() as directory:
            output=Path(directory)/'실제PDF결과.xlsx'; export(Path(__file__).resolve().parents[1]/'assets/template.xlsx',output,'SSD',products)
            wb=load_workbook(output)
            for col,p in enumerate(products,3):
                self.assertEqual(wb['SSD'].cell(6,col).value,p.name)
                self.assertEqual(wb['SSD'].cell(10,col).value,p.values['순차쓰기 속도'].text)
            wb.close()
