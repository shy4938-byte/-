"""Packaged contract/OCR verification. The HTTP server is a fixture, NOT an LLM."""
import json
import threading
from http.server import BaseHTTPRequestHandler,ThreadingHTTPServer
from pathlib import Path
import fitz
from openpyxl import load_workbook
from .ai import AIConfig,analyze_ai
from .document import read_document
from .excel import export_evidence


def verify_ai_contract(folder):
    folder=Path(folder);folder.mkdir(parents=True,exist_ok=True)
    pdf=folder/'AI 계약 검증.pdf'
    with fitz.open() as doc:
        p=doc.new_page();p.insert_text((40,60),'SSD DEMO\n1TB DEMO-ONE\n2TB DEMO-TWO\nSeries read: Up to 5500 MB/s',fontsize=16);doc.save(pdf)
        image=p.get_pixmap(matrix=fitz.Matrix(2,2)).tobytes('png')
    models=[]
    for code,cap in [('DEMO-ONE','1TB'),('DEMO-TWO','2TB')]:
        f=dict(key='용량',value=cap,scope='model',page=1,quote=cap+' '+code,link_page=1,link_quote=code,conditions='',model_ids=[code])
        models.append(dict(id=code,name=code,category='SSD',page=1,quote=cap+' '+code,fields=[f]))
    maximum=dict(key='순차읽기 속도',value='5500 MB/s',scope='family_max',page=1,quote='Series read: Up to 5500 MB/s',link_page=1,link_quote='Series read',conditions='',model_ids=['DEMO-ONE','DEMO-TWO'])
    payload=dict(models=models,shared=[maximum]);calls=[]
    class Handler(BaseHTTPRequestHandler):
        def do_POST(self):
            calls.append(json.loads(self.rfile.read(int(self.headers['Content-Length']))))
            raw=json.dumps({'message':{'content':json.dumps(payload)}}).encode();self.send_response(200);self.send_header('Content-Length',str(len(raw)));self.end_headers();self.wfile.write(raw)
        def log_message(self,*args):pass
    server=ThreadingHTTPServer(('127.0.0.1',0),Handler);thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
    try:products=analyze_ai(pdf,'SSD',AIConfig(endpoint='http://127.0.0.1:'+str(server.server_port)))
    finally:server.shutdown();server.server_close();thread.join()
    if len(products)!=2 or len(calls)!=2 or not calls[1]['messages'][1].get('images'):raise RuntimeError('AI contract transport failed')
    for p in products:
        if p.values['순차읽기 속도'].text!='확인 필요' or not p.shared:raise RuntimeError('Family maximum leaked into model specification')
    report=folder/'AI 분석근거.xlsx';export_evidence(report,products);wb=load_workbook(report)
    if '제품군 공통·미확정' not in wb.sheetnames:raise RuntimeError('Missing shared evidence workbook')
    wb.close()
    scan=folder/'OCR 스캔.pdf'
    with fitz.open() as doc:p=doc.new_page();p.insert_image(p.rect,stream=image);doc.save(scan)
    parsed=read_document(scan,ocr=True)
    if 'DEMO-ONE' not in parsed.text or '5500' not in parsed.text:raise RuntimeError('Bundled OCR failed')
    return {'provider_contract':'passed (deterministic local fixture server)','page_images':'passed','family_max_scope':'passed','evidence_excel':'passed','bundled_ocr':'passed','live_llm':'UNVERIFIED: no model or API credentials supplied'}
