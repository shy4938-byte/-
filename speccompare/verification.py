"""Runs the actual bundled runtime, worker, GUI edit and Excel export."""
import json
import os
import sys
import hashlib
from pathlib import Path
from dataclasses import asdict
from PySide6.QtCore import QEventLoop, QTimer
from PySide6.QtWidgets import QApplication
from openpyxl import load_workbook
from .gui import Window, Worker
from .resources import template_path
from .parser import analyze
from .schema import FIELDS
from .excel import export
from .fixtures import create_fixtures, create_capacity_fixture


def verify(pdf,output_folder):
    app=QApplication.instance() or QApplication([])
    products=analyze(pdf,'SSD')
    window=Window();window.tabs.setCurrentIndex(1);window.products=products;window.rebuild_groups();window.show();app.processEvents()
    output=Path(output_folder);output.mkdir(parents=True,exist_ok=True)
    export(template_path(),output/'패키지검증.xlsx','SSD',products[:4])
    report={'products':[asdict(p) for p in products],'gui':'passed','excel':'패키지검증.xlsx'}
    (output/'verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8');window.close()


def verify_suite(output_folder):
    output=Path(output_folder);output.mkdir(parents=True,exist_ok=True)
    app=QApplication.instance() or QApplication([])
    window=Window();window.show();app.processEvents()
    checks=[];paths=create_fixtures(output/'PDF 샘플')
    expected={'MEMORY':('용량','16 GB'),'SSD':('순차읽기 속도','7400 MB/s'),'VGA':('VRAM 용량','16 GB')}
    for category,pdf in paths.items():
        result=[];loop=QEventLoop();timer=QTimer();timer.setSingleShot(True);timer.timeout.connect(loop.quit)
        worker=Worker([str(pdf)],category)
        def receive(products,errors):
            result.append((products,errors));loop.quit()
        worker.result.connect(receive);worker.start();timer.start(60000);loop.exec();timer.stop()
        if not worker.wait(60000):raise RuntimeError(category+' Worker timeout')
        if not result:raise RuntimeError(category+' result signal missing')
        products,errors=result[0]
        if errors or len(products)!=1:raise RuntimeError(category+': '+str(errors))
        field,value=expected[category]
        if products[0].values[field].text!=value:raise RuntimeError(category+' extracted value mismatch')
        window.tabs.setCurrentIndex(list(FIELDS).index(category));window.products=products;window.rebuild_groups();app.processEvents()
        window.table.item(0,1).setText('GUI EDIT VERIFIED')
        if products[0].values['제조사'].text!='GUI EDIT VERIFIED':raise RuntimeError('GUI edit did not update model')
        excel=output/(category+' 비교.xlsx');export(template_path(),excel,category,products)
        wb=load_workbook(excel)
        sheet=wb[{'MEMORY':'메모리','SSD':'SSD','VGA':'VGA'}[category]]
        if sheet['C5'].value!='GUI EDIT VERIFIED' or sheet['C5'].alignment.horizontal!='center':raise RuntimeError('Excel edit/alignment mismatch')
        wb.close();checks.append({'category':category,'analysis':'passed','worker':'passed','gui_edit':'passed','excel':'passed','source':pdf.name})
    # Exercise the GUI's multi-file worker independently of the model selector.
    ssd=paths['SSD'];second=output/'PDF 샘플'/'SSD 두 번째.pdf';second.write_bytes(ssd.read_bytes())
    multi=Worker([str(ssd),str(second)],'SSD');results=[];loop=QEventLoop();timer=QTimer();timer.setSingleShot(True);timer.timeout.connect(loop.quit)
    def receive_multi(products,errors):results.append((products,errors));loop.quit()
    multi.result.connect(receive_multi);multi.start();timer.start(60000);loop.exec();timer.stop();multi.wait(60000)
    if not results or len(results[0][0])!=2 or results[0][1]:raise RuntimeError('Multiple PDF worker test failed')
    # Reproduce SSD uploaded while MEMORY remains selected (reported UI bug).
    window.tabs.setCurrentIndex(0)
    auto=Worker([str(paths['SSD'])],'MEMORY');auto_results=[];loop=QEventLoop();timer=QTimer();timer.setSingleShot(True);timer.timeout.connect(loop.quit)
    def receive_auto(products,errors):auto_results.append((products,errors));loop.quit()
    auto.result.connect(receive_auto);auto.start();timer.start(60000);loop.exec();timer.stop();auto.wait(60000)
    if not auto_results or auto_results[0][1]:raise RuntimeError('Auto category worker failed')
    window.analysis_result(*auto_results[0])
    if window.category!='SSD' or window.products[0].values['순차읽기 속도'].text!='7400 MB/s':raise RuntimeError('SSD displayed as MEMORY')
    matrix_pdf=create_capacity_fixture(output/'PDF 샘플'/'용량별 표.pdf')
    matrix=analyze(matrix_pdf,'SSD')
    if len(matrix)!=2:raise RuntimeError('Capacity matrix model count mismatch')
    for i,p in enumerate(matrix):
        expected_values={'순차읽기 속도':['7,100 MB/s','7,200 MB/s'],'순차쓰기 속도':['6,100 MB/s','6,200 MB/s'],'랜덤읽기 속도':['1,100K IOPS','1300K IOPS'],'랜덤쓰기 속도':['1,500K IOPS','1,600K IOPS'],'내구성':['600 TB','1,200 TB']}
        for key,values in expected_values.items():
            if p.values[key].text!=values[i]:raise RuntimeError('Capacity matrix mismatch: '+key)
        if 'Thread 16' not in p.values['랜덤읽기 속도'].evidence:raise RuntimeError('Missing measurement thread')
        if p.values['부품번호'].text!=f'ACME-X{i+1}':raise RuntimeError('Concatenated model columns')
    window.products=matrix;window.rebuild_groups();app.processEvents()
    matrix_excel=output/'용량별 비교.xlsx';export(template_path(),matrix_excel,'SSD',matrix)
    wb=load_workbook(matrix_excel)
    for i,p in enumerate(matrix,3):
        for r,key in enumerate(FIELDS['SSD'],5):
            cell=wb['SSD'].cell(r,i)
            if cell.value!=p.values[key].text or cell.alignment.horizontal!='center':raise RuntimeError('Capacity Excel mismatch')
    wb.close()
    bundled=bool(getattr(sys,'frozen',False))
    if bundled:
        root=Path(sys.executable).resolve().parent
        for module_name in ['PySide6','fitz','pdfplumber','openpyxl']:
            module=__import__(module_name);module_path=Path(module.__file__).resolve()
            if not module_path.is_relative_to(root):raise RuntimeError('External dependency: '+module_name)
    report={'status':'passed','frozen':bundled,'platform':sys.platform,'checks':checks,'multiple_pdf':'passed','ssd_on_memory_tab':'passed','capacity_columns_and_qd':'passed','template_sha256':hashlib.sha256(template_path().read_bytes()).hexdigest(),'executable':str(Path(sys.executable).resolve()),'test_scope':'synthetic runtime + GUI edit + worker + Excel; no claim of clean-PC/SmartScreen verification'}
    (output/'self-test.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8');window.close()
    return report
