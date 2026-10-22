from pathlib import Path
import os,tempfile
from openpyxl import Workbook
from openpyxl.styles import Alignment,Font,PatternFill
from openpyxl.utils import get_column_letter
from .analytics import monthly_rows,daily_rows,required_dates,cards
from .domain import SOURCE

RULE='USD/KRW · 최초 1회차 매매기준율 · 월평균=고시가 있는 날짜의 최초 고시값 산술평균(공식 월평균과 다를 수 있음). 휴일·누락일 값 대입 없음.'
COMPARE_RULE='상단은 오늘 또는 최근 14일 내 유효 고시일 기준 전년 동일 월일 비교. 전년 동일일 고시 없음/2월29일 대응일 없음이면 비교 불가.'

def export_snapshot(snapshot,destination):
    wb=Workbook();wb.remove(wb.active)
    def create(name,headers):
        ws=wb.create_sheet(name);ws.append([name+' · '+snapshot.asof.isoformat()]);ws.append(['출처',SOURCE]);ws.append(['조회 기준',RULE]);ws.append(['비교 규칙',COMPARE_RULE]);ws.append(['조회시각',snapshot.retrieved_at.isoformat(),'새 조회일 수',len(snapshot.fresh),'저장·누락 자료 포함 여부',True]);ws.append([]);ws.append(headers);ws.freeze_panes='B8';return ws
    monthly=create('월평균_비교',['월',f'{snapshot.asof.year} 평균',f'{snapshot.asof.year-1} 평균','차이(원)','변동률','올해 자료 상태','작년 자료 상태'])
    for row in monthly_rows(snapshot):monthly.append([row['month'],*[float(row[k]) if row[k] is not None else None for k in ['current','previous','delta']],float(row['percent']/100) if row['percent'] is not None else None,row['current_status'],row['previous_status']])
    daily=create('당월_일별_비교',['날짜',f'{snapshot.asof.year} 환율',f'{snapshot.asof.year-1} 환율','차이(원)','변동률','올해 자료 상태','작년 자료 상태'])
    for row in daily_rows(snapshot):daily.append([f'{snapshot.asof.month:02d}/{row["day"]:02d}',*[float(row[k]) if row[k] is not None else None for k in ['current','previous','delta']],float(row['percent']/100) if row['percent'] is not None else None,row['current_status'],row['previous_status']])
    raw=create('환율_원본데이터',['고시 기준일','통화','고시회차','매매기준율','고시시각','수집시각','출처','검증 상태','조회 상태','비고'])
    dates=set(required_dates(snapshot.asof));top=cards(snapshot)
    for k in ['current','previous']:
        if top[k]:dates.add(top[k].day)
    for day in sorted(dates):
        fixing=snapshot.records.get(day);coverage=snapshot.coverage.get(day)
        raw.append([day.isoformat(),'USD',1,float(fixing.rate) if fixing else None,fixing.announced_at.isoformat() if fixing else None,fixing.collected_at.isoformat() if fixing else coverage.checked_at.isoformat() if coverage else None,fixing.source if fixing else SOURCE,fixing.validation if fixing else '검증된 환율 없음',coverage.status if coverage else 'pending',coverage.detail if coverage else '미수집'])
    for ws in wb:
        ws.auto_filter.ref=f'A7:{get_column_letter(ws.max_column)}{max(7,ws.max_row)}'
        for row in ws:
            for cell in row:
                if isinstance(cell.value,str):cell.data_type='s'
                cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
        for cell in ws[7]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='087F8C')
        for i in range(1,ws.max_column+1):ws.column_dimensions[get_column_letter(i)].width=24 if i<6 else 42
    for ws in [monthly,daily]:
        for row in ws.iter_rows(min_row=8):
            for c in row[1:4]:c.number_format='#,##0.00;[Red]-#,##0.00'
            row[4].number_format='0.00%'
    for row in raw.iter_rows(min_row=8,min_col=4,max_col=4):row[0].number_format='#,##0.00'
    destination=Path(destination);temporary=None
    try:
        with tempfile.NamedTemporaryFile(dir=destination.parent,suffix='.xlsx',delete=False) as f:temporary=f.name
        wb.save(temporary);os.replace(temporary,destination);temporary=None
    finally:
        wb.close()
        if temporary and Path(temporary).exists():Path(temporary).unlink()
