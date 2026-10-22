import os
os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
import unittest,tempfile,threading
from pathlib import Path
from datetime import date,datetime,timedelta
from decimal import Decimal
from unittest.mock import Mock,patch
from PySide6.QtWidgets import QApplication
from openpyxl import load_workbook
from app.domain import ValidationError,NoPublication,AccessRestricted,Cancelled,HanaError,KST,now_utc
from app.parser import parse_fixing
from app.hana_client import HanaClient,payload_for
from app.database import Database,Coverage
from app.analytics import Snapshot,required_dates,monthly_rows,daily_rows,cards,difference
from app.service import collect
from app.dashboard import Window
from app.excel import export_snapshot
from helpers import fixing,FakeClient

SAMPLE=(Path(__file__).parent/'fixtures/sample_20261022.html').read_text(encoding='utf-8')
app=QApplication.instance() or QApplication([])

class ParserTests(unittest.TestCase):
    def test_provided_sample_not_live_rate(self):
        f=parse_fixing(SAMPLE,date(2026,10,22));self.assertEqual(f.rate,Decimal('1325.50'));self.assertEqual(f.currency,'USD');self.assertEqual(f.round,1);self.assertEqual(f.announced_at.hour,8)
        self.assertNotEqual(f.rate,Decimal('1348.69'));self.assertNotEqual(f.rate,Decimal('1312.60'))
    def test_semantic_column_not_fixed_position(self):
        html=SAMPLE[:SAMPLE.index('<table>')]+'''<table><thead><tr><th>통화</th><th>매매 기준율</th><th>현찰 살때</th></tr></thead><tbody><tr><td>일본 JPY</td><td>900.00</td><td>910.00</td></tr><tr><td>미국 USD</td><td>1,325.50</td><td>1,348.69</td></tr></tbody></table>'''
        self.assertEqual(parse_fixing(html,date(2026,10,22)).rate,Decimal('1325.50'))
    def test_wrong_currency_rejected(self):
        with self.assertRaises(ValidationError):parse_fixing(SAMPLE.replace('USD','JPY'),date(2026,10,22))
    def test_wrong_round_rejected(self):
        with self.assertRaises(ValidationError):parse_fixing(SAMPLE.replace('(1회차)','(2회차)'),date(2026,10,22))
    def test_later_date_response_rejected(self):
        with self.assertRaises(ValidationError):parse_fixing(SAMPLE,date(2026,10,21))
    def test_prior_publication_not_assigned_to_requested_day(self):
        with self.assertRaises(NoPublication):parse_fixing(SAMPLE,date(2026,10,23))
    def test_explicit_no_publication(self):
        with self.assertRaises(NoPublication):parse_fixing('<p>조회된 고시 데이터가 없습니다.</p>',date(2026,10,24))
    def test_missing_header_never_guesses_cash_rate(self):
        with self.assertRaises(ValidationError):parse_fixing(SAMPLE.replace('매매기준율','기타환율'),date(2026,10,22))
    def test_invalid_or_ambiguous_data_rejected(self):
        for html in [SAMPLE.replace('1,325.50','NaN'),SAMPLE.replace('1,325.50','0.00'),SAMPLE.replace('1,325.50','-1.00'),SAMPLE.replace('</tbody>',SAMPLE[SAMPLE.index('<tbody>')+7:SAMPLE.index('</tbody>')]+'</tbody>')]:
            with self.subTest(html=html[-50:]):
                with self.assertRaises(ValidationError):parse_fixing(html,date(2026,10,22))

class AnalyticsTests(unittest.TestCase):
    def test_october_to_november_ranges(self):
        a=Snapshot(date(2026,10,31));b=Snapshot(date(2026,11,1));self.assertEqual(len(monthly_rows(a)),9);self.assertEqual(len(monthly_rows(b)),10);self.assertEqual(len(daily_rows(a)),31);self.assertEqual(len(daily_rows(b)),1)
        self.assertIn(date(2025,10,31),required_dates(b.asof));self.assertNotIn(date(2026,11,2),required_dates(b.asof))
    def test_december_to_january_and_leap_february(self):
        self.assertEqual(len(monthly_rows(Snapshot(date(2027,1,1)))),0);self.assertEqual(len(monthly_rows(Snapshot(date(2026,12,31)))),11)
        rows=daily_rows(Snapshot(date(2024,2,29)));self.assertEqual(len(rows),29);self.assertIsNone(rows[-1]['previous']);self.assertEqual(rows[-1]['previous_status'],'날짜 없음');self.assertEqual(len(required_dates(date(2024,2,29))),119)
    def test_mean_of_first_publications_no_weekend_fill(self):
        ds=[date(2026,1,2),date(2026,1,5)];records={ds[0]:fixing(ds[0],'1300.00'),ds[1]:fixing(ds[1],'1400.00')};snap=Snapshot(date(2026,2,1),records,{d:Coverage(d,'ok',now_utc()) for d in ds})
        row=monthly_rows(snap)[0];self.assertEqual(row['current'],Decimal('1350'));self.assertFalse(row['complete'])
        daily=daily_rows(Snapshot(date(2026,1,5),records));self.assertIsNone(daily[2]['current']);self.assertIsNone(daily[3]['current'])
    def test_previous_year_missing_is_not_zero(self):
        today=date(2026,10,22);snap=Snapshot(today,{today:fixing(today)});self.assertIsNone(cards(snap)['delta']);self.assertIsNone(daily_rows(snap)[-1]['delta'])
    def test_recent_publication_has_actual_date(self):
        today=date(2026,10,25);friday=date(2026,10,23);prior=date(2025,10,23);top=cards(Snapshot(today,{friday:fixing(friday,'1400'),prior:fixing(prior,'1300')}))
        self.assertFalse(top['is_today']);self.assertEqual(top['current'].day,friday);self.assertEqual(top['comparison_date'],prior);self.assertEqual(top['delta'],Decimal('100'))
    def test_first_day_before_publication_and_leap_card(self):
        previous=date(2026,12,31);top=cards(Snapshot(date(2027,1,1),{previous:fixing(previous)}));self.assertEqual(top['current'].day,previous);self.assertEqual(top['comparison_date'],date(2025,12,31))
        day=date(2024,2,29);self.assertIsNone(cards(Snapshot(day,{day:fixing(day)}))['comparison_date'])
    def test_percent_uses_prior_rate(self):
        delta,percentage=difference(Decimal(110),Decimal(100));self.assertEqual(delta,10);self.assertEqual(percentage,10)

class CollectionTests(unittest.TestCase):
    def test_payload_dates_and_case(self):
        p=payload_for(date(2026,10,22),'tmplnqStrDt');self.assertEqual(p['tmplnqStrDt'],'2026-10-22');self.assertEqual(p['inqStrDt'],'20261022');self.assertEqual(p['curCd'],'USD');self.assertEqual(p['pbldDvCd'],'1');self.assertEqual(p['pbldSqn'],'')
    def test_today_revalidated_history_reused(self):
        with tempfile.TemporaryDirectory() as temp:
            today=date(2026,1,2);yesterday=today-timedelta(days=1);path=Path(temp)/'한글/환율.db'
            with Database(path) as db:
                db.put(fixing(yesterday));db.put(fixing(today));self.assertTrue(db.reusable(yesterday,today));self.assertFalse(db.reusable(today,today));self.assertFalse(db.reusable(yesterday,today,True));loaded,_=db.load(today);self.assertEqual(loaded[today].rate,fixing(today).rate)
    def test_cancel_commits_partial_and_does_not_erase_cache(self):
        with tempfile.TemporaryDirectory() as temp:
            day=date(2026,1,2);event=threading.Event();client=FakeClient(records={day:fixing(day)});path=Path(temp)/'data.db'
            def update(done,total,message):
                if done==1:event.set()
            result=collect(path,day,event,progress=update,client_factory=lambda **kwargs:client);self.assertTrue(result.cancelled);self.assertIn(day,result.records);self.assertTrue(client.closed)
    def test_offline_cache_preserved_and_error_marked(self):
        class Offline(FakeClient):
            def fetch(self,day):raise AccessRestricted('blocked')
        with tempfile.TemporaryDirectory() as temp:
            day=date(2026,1,2);path=Path(temp)/'data.db'
            with Database(path) as db:db.put(fixing(day))
            result=collect(path,day,client_factory=Offline);self.assertIn(day,result.records);self.assertTrue(result.error);self.assertEqual(result.fresh,set())
    def test_access_restriction_not_retried_or_bypassed(self):
        session=Mock();session.request.return_value=Mock(status_code=403);client=HanaClient(session=session)
        with self.assertRaises(AccessRestricted):client._request('GET','https://www.kebhana.com/robots.txt')
        self.assertEqual(session.request.call_count,1)
    def test_robots_disallow_stops_before_query(self):
        session=Mock();session.request.return_value=Mock(status_code=200,text='User-agent: *\nDisallow: /cms/',content=b'')
        with self.assertRaises(AccessRestricted):HanaClient(session=session).initialize()
        self.assertEqual(session.request.call_count,1)
    def test_retry_bounded_wait_cancellable(self):
        session=Mock();session.request.return_value=Mock(status_code=503,headers={});client=HanaClient(session=session)
        with patch.object(client,'_wait'):
            with self.assertRaises(HanaError):client._request('GET','https://www.kebhana.com/robots.txt')
        self.assertEqual(session.request.call_count,3);client.cancel.set()
        with self.assertRaises(Cancelled):client._wait(1)

class ExportAndGuiTests(unittest.TestCase):
    def test_excel_sheets_metadata_numeric_format(self):
        with tempfile.TemporaryDirectory() as temp:
            day=date(2026,10,22);snap=Snapshot(day,{day:fixing(day)});target=Path(temp)/'환율.xlsx';export_snapshot(snap,target)
            wb=load_workbook(target);self.assertEqual(wb.sheetnames,['월평균_비교','당월_일별_비교','환율_원본데이터']);self.assertIn('kebhana.com',wb.worksheets[0]['B2'].value);self.assertEqual(wb.worksheets[1]['B29'].value,1325.5);self.assertEqual(wb.worksheets[1]['B29'].number_format,'#,##0.00;[Red]-#,##0.00');wb.close()
    def test_tables_charts_month_rollover_same_snapshot(self):
        with tempfile.TemporaryDirectory() as temp:
            clock=[datetime(2026,10,31,23,59,tzinfo=KST)];window=Window(Path(temp)/'data.db',clock=lambda:clock[0],auto_start=False);day=date(2026,10,22);window.apply_snapshot(Snapshot(clock[0].date(),{day:fixing(day)}))
            self.assertEqual(window.day_table.rowCount(),31);self.assertEqual(window.month_table.rowCount(),9);self.assertEqual(window.day_chart.plotted['current'],[(22,fixing(day).rate)]);self.assertIn('1,325.50',window.day_table.item(21,1).text())
            clock[0]=datetime(2026,11,1,0,1,tzinfo=KST);window.check_date(start=False);self.assertEqual(window.month_table.rowCount(),10);self.assertEqual(window.day_table.rowCount(),1);self.assertIn('11월',window.day_title.text());self.assertEqual(window.snapshot.asof,date(2026,11,1));window.close()
    def test_old_snapshot_cannot_overwrite_new_month(self):
        with tempfile.TemporaryDirectory() as temp:
            window=Window(Path(temp)/'data.db',clock=lambda:datetime(2027,1,1,tzinfo=KST),auto_start=False);window.apply_snapshot(Snapshot(date(2026,12,31)));self.assertEqual(window.snapshot.asof,date(2027,1,1));self.assertEqual(window.month_table.rowCount(),0);window.close()

class CapturedRegressionTests(unittest.TestCase):
    def test_real_capture_has_query_timestamp_distinct_from_publication(self):
        for day,rate in [(date(2026,10,22),'1325.50'),(date(2025,10,22),'1423.70')]:
            fragment=(Path(__file__).parent/'fixtures'/('captured_'+day.strftime('%Y%m%d')+'.html')).read_text(encoding='utf-8')
            record=parse_fixing(fragment,day);self.assertEqual(record.rate,Decimal(rate));self.assertEqual(record.announced_at.year,day.year);self.assertEqual(record.round,1)

    def test_real_weekend_and_holiday_never_copy_prior_rate(self):
        for day in [date(2026,10,18),date(2026,1,1)]:
            html=(Path(__file__).parent/'fixtures'/('captured_'+day.strftime('%Y%m%d')+'.html')).read_text(encoding='utf-8')
            with self.assertRaises(NoPublication):parse_fixing(html,day)
    def test_conflicting_no_data_preserves_cache_but_is_not_complete(self):
        with tempfile.TemporaryDirectory() as temp:
            day=date(2026,1,2)
            with Database(Path(temp)/'data.db') as db:
                db.put(fixing(day));db.mark(day,'no_data','no data')
                records,coverage=db.load(day);self.assertIn(day,records);self.assertEqual(coverage[day].status,'error')
